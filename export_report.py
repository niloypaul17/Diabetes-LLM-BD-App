"""
Export the Prediction Explanation result (patient inputs + risk estimate +
top contributing factors) as PDF, DOCX, PNG, or JPEG.
"""

import io
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


def _report_lines(patient_input, probability, threshold, prediction_text, top_factors):
    lines = [
        "DiaLLM-BD — Prediction Explanation Report",
        "Research prototype only — not a medical diagnosis.",
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        "",
        f"Result: {prediction_text}",
        f"Estimated probability: {probability:.3f}",
        f"Decision threshold: {threshold:.2f}",
        "",
        "Which values influenced this estimate:",
    ]
    for f in top_factors:
        sign = "+" if f["contribution"] >= 0 else ""
        lines.append(f"  - {f['feature']} (value: {f['value']}): {sign}{f['contribution']:.3f}")
    lines += ["", "Entered patient information:"]
    for k, v in patient_input.items():
        lines.append(f"  - {k}: {v}")
    lines += [
        "",
        "This is an AI-based decision-support output only. A healthcare professional",
        "should confirm status using fasting glucose, HbA1c, or an oral glucose",
        "tolerance test.",
    ]
    return lines


def make_waterfall_figure(top_factors, probability, threshold):
    fig, ax = plt.subplots(figsize=(5, 2.8))
    features = [f["feature"] for f in top_factors][::-1]
    values = [f["contribution"] for f in top_factors][::-1]
    colors = ["#d62728" if v >= 0 else "#1f77b4" for v in values]
    ax.barh(features, values, color=colors)
    ax.axvline(0, color="black", linewidth=0.8)
    ax.set_xlabel("Contribution to risk estimate (SHAP value)")
    ax.set_title(f"Probability {probability:.3f} vs threshold {threshold:.2f}")
    for i, v in enumerate(values):
        ax.text(v, i, f"{v:+.3f}", va="center", ha="left" if v >= 0 else "right", fontsize=8)
    fig.tight_layout()
    return fig


def _load_font(size):
    from PIL import ImageFont
    for path in ("arial.ttf", "Arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(path, size)
        except Exception:
            continue
    return ImageFont.load_default()


def export_png_or_jpeg(patient_input, probability, threshold, prediction_text, top_factors, fmt="png"):
    """Full report (text + chart) as one image, so PNG/JPEG carry the same
    information as the PDF/DOCX exports, not just the chart."""
    from PIL import Image, ImageDraw

    lines = _report_lines(patient_input, probability, threshold, prediction_text, top_factors)

    fig = make_waterfall_figure(top_factors, probability, threshold)
    chart_buf = io.BytesIO()
    fig.savefig(chart_buf, format="png", dpi=150)
    plt.close(fig)
    chart_buf.seek(0)
    chart_img = Image.open(chart_buf).convert("RGB")

    font = _load_font(15)
    title_font = _load_font(19)
    line_height = 22
    text_width = 760
    text_height = line_height * (len(lines) + 2)

    text_img = Image.new("RGB", (max(text_width, chart_img.width), text_height), "white")
    draw = ImageDraw.Draw(text_img)
    y = 12
    for i, line in enumerate(lines):
        draw.text((24, y), line, fill="black", font=title_font if i == 0 else font)
        y += line_height + (6 if i == 0 else 0)

    total_width = max(text_img.width, chart_img.width)
    total_height = text_img.height + chart_img.height + 20
    canvas_img = Image.new("RGB", (total_width, total_height), "white")
    canvas_img.paste(text_img, (0, 0))
    canvas_img.paste(chart_img, ((total_width - chart_img.width) // 2, text_img.height + 10))

    buf = io.BytesIO()
    if fmt == "jpeg":
        canvas_img.save(buf, format="JPEG", quality=92)
    else:
        canvas_img.save(buf, format="PNG")
    buf.seek(0)
    return buf.getvalue()


def export_pdf(patient_input, probability, threshold, prediction_text, top_factors):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import cm

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    y = height - 2 * cm

    lines = _report_lines(patient_input, probability, threshold, prediction_text, top_factors)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(2 * cm, y, lines[0])
    y -= 0.8 * cm
    c.setFont("Helvetica", 10)
    for line in lines[1:]:
        if y < 2 * cm:
            c.showPage()
            y = height - 2 * cm
            c.setFont("Helvetica", 10)
        c.drawString(2 * cm, y, line[:110])
        y -= 0.55 * cm

    # embed the waterfall chart on a fresh page
    fig = make_waterfall_figure(top_factors, probability, threshold)
    img_buf = io.BytesIO()
    fig.savefig(img_buf, format="png", dpi=200)
    plt.close(fig)
    img_buf.seek(0)
    from reportlab.lib.utils import ImageReader
    c.showPage()
    c.drawImage(ImageReader(img_buf), 2 * cm, height / 2 - 6 * cm, width=width - 4 * cm, preserveAspectRatio=True)

    c.save()
    buf.seek(0)
    return buf.getvalue()


def export_docx(patient_input, probability, threshold, prediction_text, top_factors):
    from docx import Document
    from docx.shared import Pt, Inches

    doc = Document()
    doc.add_heading("DiaLLM-BD — Prediction Explanation Report", level=1)
    p = doc.add_paragraph("Research prototype only — not a medical diagnosis.")
    p.runs[0].italic = True
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    doc.add_heading("Result", level=2)
    doc.add_paragraph(f"{prediction_text}")
    doc.add_paragraph(f"Estimated probability: {probability:.3f}")
    doc.add_paragraph(f"Decision threshold: {threshold:.2f}")

    doc.add_heading("Which values influenced this estimate", level=2)
    for f in top_factors:
        sign = "+" if f["contribution"] >= 0 else ""
        doc.add_paragraph(f"{f['feature']} (value: {f['value']}): {sign}{f['contribution']:.3f}", style="List Bullet")

    fig = make_waterfall_figure(top_factors, probability, threshold)
    img_buf = io.BytesIO()
    fig.savefig(img_buf, format="png", dpi=200)
    plt.close(fig)
    img_buf.seek(0)
    doc.add_picture(img_buf, width=Inches(6))

    doc.add_heading("Entered patient information", level=2)
    for k, v in patient_input.items():
        doc.add_paragraph(f"{k}: {v}", style="List Bullet")

    doc.add_heading("Note", level=2)
    note = doc.add_paragraph(
        "This is an AI-based decision-support output only. A healthcare professional "
        "should confirm status using fasting glucose, HbA1c, or an oral glucose "
        "tolerance test."
    )
    for run in note.runs:
        run.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()
